from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlmodel import Session, select

from backend.app.models import AuditLog, InventoryLocation, InventoryMovement
from backend.app.services.auth_service import (
    CSRF_COOKIE,
    CSRF_HEADER,
    SESSION_COOKIE,
    create_session,
    create_user,
)
from backend.app.services.inventory_excel import (
    parse_inventory_file_rows,
    render_reconcile_export_xlsx,
)

HEADERS = [
    "factory_id",
    "part_key",
    "location_code",
    "quantity",
    "excel_quantity",
    "delta",
    "category",
    "note",
]


def _login(client: TestClient, session: Session, *, username: str, role: str) -> None:
    user = create_user(
        session,
        username=username,
        display_name=username.replace("-", " ").title(),
        password=f"{username}-pass",
        role=role,
    )
    token, _ = create_session(session, user, ip_address="testclient", user_agent="pytest")
    client.cookies.set(SESSION_COOKIE, token)
    client.cookies.set(CSRF_COOKIE, "pytest-csrf-token")
    client.headers.update({CSRF_HEADER: "pytest-csrf-token"})


def _seed_location(
    session: Session,
    *,
    part_key: str,
    location_code: str,
    quantity: int,
    factory_id: str = "factory_a",
) -> InventoryLocation:
    row = InventoryLocation(
        factory_id=factory_id,
        part_key=part_key,
        location_code=location_code,
        quantity=quantity,
        status="active" if quantity > 0 else "zero_stock",
        zero_stock=quantity <= 0,
        location_kind="permanent",
    )
    session.add(row)
    session.commit()
    session.refresh(row)
    return row


def _xlsx_upload(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["factory_id", "part_key", "location_code", "quantity"])
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _sheet_rows(content: bytes) -> list[tuple[object, ...]]:
    workbook = load_workbook(BytesIO(content), data_only=True)
    return list(workbook.active.iter_rows(values_only=True))


def test_render_reconcile_export_xlsx_columns_values_order_and_reparse() -> None:
    content = render_reconcile_export_xlsx(
        {
            "matched": [
                {
                    "factory_id": "factory_a",
                    "part_key": "CPXS000122400",
                    "location_code": "A-A01-010",
                    "system_quantity": 4,
                    "excel_quantity": 4,
                }
            ],
            "quantity_mismatch": [
                {
                    "factory_id": "factory_a",
                    "part_key": "CPXS000122200",
                    "location_code": "A-A01-020",
                    "system_quantity": 5,
                    "excel_quantity": 7,
                    "delta": 2,
                }
            ],
            "excel_missing": [
                {
                    "factory_id": "factory_a",
                    "part_key": "CPXS000122100",
                    "location_code": "A-A01-030",
                    "system_quantity": 3,
                }
            ],
            "excel_new": [
                {
                    "factory_id": "factory_a",
                    "part_key": "CPXS000122300",
                    "location_code": "TMP-01",
                    "excel_quantity": 9,
                }
            ],
            "summary": {},
        }
    )

    rows = _sheet_rows(content)

    assert list(rows[0]) == HEADERS
    assert [row[6] for row in rows[1:]] == [
        "quantity_mismatch",
        "excel_missing",
        "excel_new",
        "matched",
    ]
    assert rows[1][:7] == (
        "factory_a",
        "CPXS000122200",
        "A-A01-020",
        5,
        7,
        2,
        "quantity_mismatch",
    )
    assert rows[2][3] == 3
    assert rows[2][4] is None
    assert rows[2][5] is None
    assert rows[3][3] == 0
    assert rows[3][4] == 9
    assert rows[4][3] == 4
    assert rows[4][4] == 4
    assert all(isinstance(row[7], str) and row[7] for row in rows[1:])

    reparsed = parse_inventory_file_rows(filename="reconcile-export.xlsx", content=content)
    assert reparsed == [
        {
            "factory_id": "factory_a",
            "part_key": "CPXS000122200",
            "location_code": "A-A01-020",
            "quantity": 5,
        },
        {
            "factory_id": "factory_a",
            "part_key": "CPXS000122100",
            "location_code": "A-A01-030",
            "quantity": 3,
        },
        {
            "factory_id": "factory_a",
            "part_key": "CPXS000122300",
            "location_code": "TMP-01",
            "quantity": 0,
        },
        {
            "factory_id": "factory_a",
            "part_key": "CPXS000122400",
            "location_code": "A-A01-010",
            "quantity": 4,
        },
    ]


def test_render_reconcile_export_xlsx_escapes_formula_prefixes_and_sorts_within_groups() -> None:
    content = render_reconcile_export_xlsx(
        {
            "matched": [],
            "quantity_mismatch": [
                {
                    "factory_id": "=factory",
                    "part_key": "+PART-B",
                    "location_code": "@LOC-B",
                    "system_quantity": 1,
                    "excel_quantity": 2,
                    "delta": 1,
                },
                {
                    "factory_id": "factory_a",
                    "part_key": "-PART-C",
                    "location_code": "=LOC-C",
                    "system_quantity": 2,
                    "excel_quantity": 1,
                    "delta": -1,
                },
                {
                    "factory_id": "factory_a",
                    "part_key": "=PART-A",
                    "location_code": "A-A01-001",
                    "system_quantity": 2,
                    "excel_quantity": 1,
                    "delta": -1,
                },
                {
                    "factory_id": "factory_a",
                    "part_key": "@PART-D",
                    "location_code": "A-A01-004",
                    "system_quantity": 2,
                    "excel_quantity": 1,
                    "delta": -1,
                },
            ],
            "excel_missing": [],
            "excel_new": [],
        }
    )

    rows = _sheet_rows(content)

    assert [row[1] for row in rows[1:]] == [
        "'+PART-B",
        "'-PART-C",
        "'=PART-A",
        "'@PART-D",
    ]
    assert rows[1][0] == "'=factory"
    assert rows[1][2] == "'@LOC-B"
    assert rows[2][2] == "'=LOC-C"


def test_reconcile_export_file_endpoint_requires_supervisor(
    client: TestClient,
    session: Session,
) -> None:
    _login(client, session, username="reconcile-export-operator", role="operator")
    upload = _xlsx_upload([["factory_a", "CPXS000122501", "A-A01-001", 1]])

    response = client.post(
        "/api/inventory/reconcile/export-file",
        files={
            "file": (
                "stocktake.xlsx",
                upload,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 403


def test_reconcile_export_file_endpoint_returns_xlsx_and_is_read_only(
    client: TestClient,
    session: Session,
) -> None:
    _login(client, session, username="reconcile-export-supervisor", role="supervisor")
    _seed_location(
        session,
        part_key="CPXS000122601",
        location_code="A-A01-001",
        quantity=5,
    )
    _seed_location(
        session,
        part_key="CPXS000122602",
        location_code="A-A01-002",
        quantity=3,
    )
    upload = _xlsx_upload(
        [
            ["factory_a", "CPXS000122601", "A-A01-001", 7],
            ["factory_a", "CPXS000122603", "TMP-03", 2],
        ]
    )
    counts_before = (
        len(session.exec(select(InventoryLocation)).all()),
        len(session.exec(select(InventoryMovement)).all()),
        len(session.exec(select(AuditLog)).all()),
    )

    response = client.post(
        "/api/inventory/reconcile/export-file",
        files={
            "file": (
                "stocktake.xlsx",
                upload,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    disposition = response.headers["content-disposition"]
    assert "attachment" in disposition
    assert "tagledger-reconcile-export-" in disposition
    assert disposition.endswith(".xlsx")
    rows = _sheet_rows(response.content)
    assert list(rows[0]) == HEADERS
    assert [row[6] for row in rows[1:]] == [
        "quantity_mismatch",
        "excel_missing",
        "excel_new",
    ]
    assert rows[1][1:6] == ("CPXS000122601", "A-A01-001", 5, 7, 2)
    assert rows[2][1:6] == ("CPXS000122602", "A-A01-002", 3, None, None)
    assert rows[3][1:6] == ("CPXS000122603", "TMP-03", 0, 2, None)
    assert parse_inventory_file_rows(filename="export.xlsx", content=response.content)

    counts_after = (
        len(session.exec(select(InventoryLocation)).all()),
        len(session.exec(select(InventoryMovement)).all()),
        len(session.exec(select(AuditLog)).all()),
    )
    assert counts_after == counts_before
