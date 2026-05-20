import re
from pathlib import Path
from typing import NoReturn

from openpyxl import load_workbook

from backend.app.config import get_settings

from .normalize import (
    OutboundItem,
    normalize_order_no,
    normalize_part_code,
)


def _read_text(path: Path) -> str:
    if not path.exists():
        return ""
    return path.read_text(encoding="utf-8", errors="ignore")


def _cell_text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _cell_int(value: object) -> int | None:
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    text = _cell_text(value)
    return int(text) if text.isdigit() else None


def _has_visible_mark(value: object) -> bool:
    text = _cell_text(value).upper()
    return text in {"X", "×", "√", "Y", "YES", "OK", "DONE", "已剪"}


def _is_invalid_location(value: str) -> bool:
    text = value.strip().upper()
    return text in {"#VALUE!", "#N/A", "#REF!", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!"}


def _row_locations(row: tuple[object, ...], *, start_index: int = 3) -> tuple[str, ...]:
    locations = []
    for cell in row[start_index:]:
        text = _cell_text(cell)
        if text and not _has_visible_mark(text) and not _is_invalid_location(text):
            locations.append(text)
    return tuple(locations)


def _normalize_header_text(value: object) -> str:
    text = _cell_text(value)
    text = text.replace("（", "(").replace("）", ")")
    return re.sub(r"\s+", "", text).lower()


def _resolve_column_indexes(sheet) -> dict[str, int]:
    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(header):
        token = _normalize_header_text(cell)
        if not token:
            continue
        if "出库单号" in token and "order_no" not in mapping:
            mapping["order_no"] = idx
        if ("数量" in token or "qty" in token or "quantity" in token) and "quantity" not in mapping:
            mapping["quantity"] = idx
        if (
            "备件编码" in token or "编码" in token or "part" in token or "sku" in token
        ) and "part_code" not in mapping:
            mapping["part_code"] = idx
        if ("备件名称" in token or "名称" in token or "name" in token) and "name" not in mapping:
            mapping["name"] = idx
        if (
            "库位" in token or "location" in token or "loc" in token
        ) and "location_start" not in mapping:
            mapping["location_start"] = idx
    return mapping


def _load_shipping_sheet(path: Path, sheet_name: str) -> list[OutboundItem]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            return []
        sheet = workbook[sheet_name]
        items: list[OutboundItem] = []
        col = _resolve_column_indexes(sheet)
        order_idx = col.get("order_no", 0)
        qty_idx = col.get("quantity", 1)
        part_idx = col.get("part_code", 2)
        name_idx = col.get("name", 3)
        current_order_no = ""
        for row in sheet.iter_rows(min_row=2, values_only=True):
            order_no = normalize_order_no(
                _cell_text(row[order_idx] if len(row) > order_idx else "")
            )
            if order_no:
                current_order_no = order_no
            else:
                order_no = current_order_no
            quantity = _cell_int(row[qty_idx] if len(row) > qty_idx else None)
            part_code = normalize_part_code(
                _cell_text(row[part_idx] if len(row) > part_idx else "")
            )
            name = _cell_text(row[name_idx] if len(row) > name_idx else "")
            if order_no and part_code:
                items.append(
                    OutboundItem(
                        order_no=order_no,
                        part_code=part_code,
                        quantity=quantity,
                        source="shipping",
                        raw_line=f"{order_no} {quantity or ''} {part_code} {name}".strip(),
                        name=name,
                    )
                )
        return items
    finally:
        workbook.close()


def _load_cutting_sheet(path: Path, sheet_name: str) -> list[OutboundItem]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            return []
        sheet = workbook[sheet_name]
        items: list[OutboundItem] = []
        col = _resolve_column_indexes(sheet)
        order_idx = col.get("order_no")
        qty_idx = col.get("quantity", 0)
        part_idx = col.get("part_code", 1)
        name_idx = col.get("name", 2)
        location_start = col.get("location_start", 3)
        for row in sheet.iter_rows(min_row=2, values_only=True):
            order_no = (
                normalize_order_no(_cell_text(row[order_idx] if len(row) > order_idx else ""))
                if order_idx is not None
                else ""
            )
            quantity = _cell_int(row[qty_idx] if len(row) > qty_idx else None)
            part_code = normalize_part_code(
                _cell_text(row[part_idx] if len(row) > part_idx else "")
            )
            name = _cell_text(row[name_idx] if len(row) > name_idx else "")
            checked = any(_has_visible_mark(cell) for cell in row[location_start:])
            locations = _row_locations(tuple(row), start_index=location_start)
            if part_code and ((quantity is not None and quantity > 0) or checked):
                items.append(
                    OutboundItem(
                        order_no=order_no or "PICKING_TOTAL",
                        part_code=part_code,
                        quantity=quantity,
                        source="cutting",
                        raw_line=f"{quantity or ''} {part_code} {name}".strip(),
                        name=name,
                        locations=locations,
                    )
                )
        return items
    finally:
        workbook.close()


def _real_order_numbers(items: list[OutboundItem]) -> set[str]:
    return {item.order_no for item in items if item.order_no and item.order_no != "PICKING_TOTAL"}


def _shipping_items_from_ordered_cutting(items: list[OutboundItem]) -> list[OutboundItem]:
    return [
        OutboundItem(
            order_no=item.order_no,
            part_code=item.part_code,
            quantity=item.quantity,
            source="shipping",
            raw_line=item.raw_line,
            name=item.name,
            locations=item.locations,
        )
        for item in items
        if item.order_no != "PICKING_TOTAL"
    ]


def _fail_outbound_config(message: str) -> NoReturn:
    raise RuntimeError(message)


def load_outbound_items() -> tuple[list[OutboundItem], list[OutboundItem]]:
    settings = get_settings()
    if not settings.outbound_workbook_file.exists():
        _fail_outbound_config(f"outbound workbook not found: {settings.outbound_workbook_file}")
    cutting = _load_cutting_sheet(settings.outbound_workbook_file, settings.outbound_cutting_sheet)
    shipping = _load_shipping_sheet(
        settings.outbound_workbook_file, settings.outbound_shipping_sheet
    )
    cutting_orders = _real_order_numbers(cutting)
    if cutting_orders:
        scoped_shipping = [
            item for item in shipping if normalize_order_no(item.order_no) in cutting_orders
        ]
        shipping = scoped_shipping or _shipping_items_from_ordered_cutting(cutting)
    if not cutting or not shipping:
        _fail_outbound_config(
            f"outbound workbook missing required rows or sheets: {settings.outbound_workbook_file}"
        )
    return cutting, shipping
